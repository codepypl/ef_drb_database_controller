from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

import config
from libs.msal_client import MailAttachment, MailMessage

COL_DELIVERY_TIME = "Godzina dostarczenia maila"
COL_SENDER = "Nadawca maila"
COL_SUBJECT = "Temat maila"
COL_FILENAME = "Nazwa pliku"
COL_SCAN_DATE = "Data skanowania"
COL_SCAN_TIME = "Godzina skanowania"

COLUMNS = (
    COL_DELIVERY_TIME,
    COL_SENDER,
    COL_SUBJECT,
    COL_FILENAME,
    COL_SCAN_DATE,
    COL_SCAN_TIME,
)


@dataclass(frozen=True)
class RegistryEntry:
    delivery_time: time
    sender: str
    subject: str
    filename: str
    scan_date: date
    scan_time: time


class Registry:
    def __init__(self, path: Path | None = None) -> None:
        self._path = path or config.REGISTRY_LOCAL_PATH
        self._entries: list[RegistryEntry] = []

    @property
    def path(self) -> Path:
        return self._path

    def load(self) -> None:
        if not self._path.exists():
            self._entries = []
            return

        workbook = load_workbook(self._path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                self._entries = []
                return
            self._entries = self._read_rows(sheet)
        finally:
            workbook.close()

    def save(self) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Rejestr"
        sheet.append(list(COLUMNS))

        for entry in self._entries:
            sheet.append(
                [
                    entry.delivery_time.strftime("%H:%M:%S"),
                    entry.sender,
                    entry.subject,
                    entry.filename,
                    entry.scan_date.strftime("%Y-%m-%d"),
                    entry.scan_time.strftime("%H:%M:%S"),
                ]
            )

        workbook.save(self._path)
        workbook.close()

    def find_by_filename(self, filename: str) -> RegistryEntry | None:
        normalized = filename.strip()
        for entry in self._entries:
            if entry.filename == normalized:
                return entry
        return None

    def add_entry(self, entry: RegistryEntry) -> None:
        self._entries.append(entry)

    @staticmethod
    def build_entry(
        message: MailMessage,
        attachment: MailAttachment,
        scanned_at: datetime,
    ) -> RegistryEntry:
        return RegistryEntry(
            delivery_time=message.received_at.time().replace(microsecond=0),
            sender=message.sender,
            subject=message.subject,
            filename=attachment.name.strip(),
            scan_date=scanned_at.date(),
            scan_time=scanned_at.time().replace(microsecond=0),
        )

    def _read_rows(self, sheet: Worksheet) -> list[RegistryEntry]:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []

        column_index = {name: index for index, name in enumerate(header)}
        missing = [column for column in COLUMNS if column not in column_index]
        if missing:
            raise ValueError(f"Registry file is missing columns: {', '.join(missing)}")

        entries: list[RegistryEntry] = []
        for row in rows:
            if row is None or all(value in (None, "") for value in row):
                continue
            entries.append(
                RegistryEntry(
                    delivery_time=_parse_time(row[column_index[COL_DELIVERY_TIME]]),
                    sender=_as_text(row[column_index[COL_SENDER]]),
                    subject=_as_text(row[column_index[COL_SUBJECT]]),
                    filename=_as_text(row[column_index[COL_FILENAME]]),
                    scan_date=_parse_date(row[column_index[COL_SCAN_DATE]]),
                    scan_time=_parse_time(row[column_index[COL_SCAN_TIME]]),
                )
            )
        return entries


def _as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_time(value: object) -> time:
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    text = _as_text(value)
    if not text:
        raise ValueError("Missing time value in registry row")
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    parsed = datetime.fromisoformat(text)
    return parsed.time().replace(microsecond=0)


def _parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _as_text(value)
    if not text:
        raise ValueError("Missing date value in registry row")
    return datetime.strptime(text, "%Y-%m-%d").date()
